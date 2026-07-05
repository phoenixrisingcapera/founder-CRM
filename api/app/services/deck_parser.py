from pathlib import Path
from io import BytesIO

import fitz
from pptx import Presentation


def parse_deck_file(file_name: str, content: bytes) -> list[dict[str, str]]:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(content)
    if suffix == ".pptx":
        return _parse_pptx(content)

    text = content.decode("utf-8", errors="ignore")
    return _slides_from_text(text)


def _parse_pdf(content: bytes) -> list[dict[str, str]]:
    document = fitz.open(stream=content, filetype="pdf")
    slides: list[dict[str, str]] = []
    for index, page in enumerate(document, start=1):
        page_text = page.get_text("text").strip()
        title = page_text.splitlines()[0] if page_text else f"Slide {index}"
        slides.append({"title": title[:255], "content": page_text or "No extractable content."})
    return slides or [{"title": "Slide 1", "content": "No extractable content."}]


def _parse_pptx(content: bytes) -> list[dict[str, str]]:
    presentation = Presentation(BytesIO(content))
    slides: list[dict[str, str]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                texts.append(text.strip())
        title = texts[0].splitlines()[0] if texts else f"Slide {index}"
        body = "\n\n".join(texts) if texts else "No extractable content."
        slides.append({"title": title[:255], "content": body})
    return slides or [{"title": "Slide 1", "content": "No extractable content."}]


def _slides_from_text(text: str) -> list[dict[str, str]]:
    chunks = [chunk.strip() for chunk in text.split("\n\n") if chunk.strip()]
    slides: list[dict[str, str]] = []
    for index, chunk in enumerate(chunks or [text.strip() or "Empty deck upload."], start=1):
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        title = lines[0] if lines else f"Slide {index}"
        slides.append({"title": title[:255], "content": chunk})
    return slides
